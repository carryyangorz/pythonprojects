import requests
import os
import time
import re
import json
import random
from openpyxl import Workbook
from openpyxl import load_workbook
from selenium import webdriver
from selenium.webdriver import ChromeOptions
from time import sleep
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait 
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.desired_capabilities import DesiredCapabilities
from threading import Timer
import threading
capa = DesiredCapabilities.CHROME
capa["pageLoadStrategy"] = "eager"
option=ChromeOptions()
# option.add_argument('--headless')
option.add_argument('--no-sandbox')
option.add_argument('log-level=3') #INFO = 0 WARNING = 1 LOG_ERROR = 2 LOG_FATAL = 3 default is 0
option.add_experimental_option('excludeSwitches',['enable-automation'])
# option.add_argument('--blink-settings=imagesEnabled=false')
proxy = [{
        'http': '60.188.5.211:80',
},
{
        'http': '114.231.41.213:8888',
},
{
        'http': '162.241.166.68:80',
},
{
        'http': '117.69.232.19:8089',
    }
]
# option.add_experimental_option("debuggerAddress", "127.0.0.1:9527")
import multiprocessing
from multiprocessing import Process,Lock
# import multiprocessing_win
headers={
        'user-agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_14_6) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/78.0.3904.97 Safari/537.36'
        }

specie_type_ls=['photo',
'audio',
'video'
]
def savedata(name):
    try:
        wb=load_workbook(name+'//' +name+'.xlsx')
    except:
        return
    ws=wb.active
    file=os.listdir(name+'//')
    for item in file:
        if not '_' in item:
            continue
        if '.txt' in item:
            print(item+'  removed')
            f=open(name+'//' +item,'r',encoding='utf-8')
            ws.append(f.readlines())
            f.close()
            wb.save(name+'//' +name+ '.xlsx')
            os.remove(name+'//' +item)
    wb.close()
def savedata_2(data):
    name=data[0][0]
    specie_type=data[0][1]
    data.remove(data[0])
    wb=load_workbook(name+"//"+name+"_"+specie_type+'.xlsx')
    ws=wb.active
    for item in data:
        ws.append(item)
    wb.save(name+"//"+name+"_"+specie_type+'.xlsx')
def download_species(name,num,type):
    if os.path.exists(name+'//'):
        print(name+' exists, pass...')
        return
    print(name+' '+type+' searching...')
    itemnum=0
    if not os.path.exists(name+'//'):
       os.makedirs(name+'//')

    url='https://search.macaulaylibrary.org/catalog?mediaType='+type+'&sort=rating_rank_desc'
    browser.get(url)
    browser.implicitly_wait(5)
    browser.find_element(By.CSS_SELECTOR,'input.Suggest-input').send_keys(name)
    browser.implicitly_wait(5)

    try:
        browser.find_element(By.CSS_SELECTOR,'div#Suggest-suggestion-0').click()
    except:
        print(name+' no result')
        return
    flag=1
    while(flag):
        js="var q=document.documentElement.scrollTop=100000"
        browser.execute_script(js)
        try:
            a=browser.find_element(By.CSS_SELECTOR,'div.pagination')
            a.find_element(By.CSS_SELECTOR,'button.u-margin-none').click()
        except:
            pass
        sleep(3)
        a=browser.find_elements(By.CSS_SELECTOR,'a.ResultsGallery-link')
        if len(a)==itemnum:
             num=itemnum
             break
        itemnum=len(a)
        print(itemnum)
        if len(a)>num:
              flag=0

    print('num'+str(num))
    if num==0:
        return

    aa=browser.find_element(By.CSS_SELECTOR,'#content > div > div > form > div.toolbar > div.filters > div.mediaType')
    if not os.path.exists(name+'//'+name+'.txt'):
        f=open(name+'//'+name+'.txt','w',encoding='utf-8')
        f.write(aa.text.replace('Remove media type filter','Filter by photo'))
        f.close()
    if not os.path.exists(name+'//' +name+'.xlsx'):
        wb=Workbook()
        ws=wb.active
        ws.append(['File Name','URL','ID','Name','SCIname','Auth','AgeSex','Date','Locate','Coordinates'])
        wb.save(name+'//' +name+'.xlsx')
    ppls=[]
    i=0
    flag=1

    while flag:
        for j in range(20):
                sleep(random.randrange(4))
                print(i)
                if i>num-1:
                    flag=0
                    break
                id=a[i].get_attribute('data-asset-id')
                data=[id,name,type,i]
                pp=Process(target=download,args=(data,))
                i=i+1
                # if i>len(bb):
                #         break
                pp.start()
        ppls.append(pp)
        for thread in ppls:
                thread.join()
        savedata(name)
    savedata(name)
    wb.close()
        
def download(data):
    option.add_argument('--headless')
    picid=data[0]
    name=data[1]
    type=data[2]
    i=str(data[3])
    url='https://macaulaylibrary.org/asset/'+picid
    if type=='photo':
        tail='.jpg'
        picurl='https://cdn.download.ams.birds.cornell.edu/api/v2/asset/'+picid +'/2400'
    elif type=='audio':
        tail='.mp3'
        picurl='https://cdn.download.ams.birds.cornell.edu/api/v2/asset/'+picid +'/mp3'
    else:
        tail='.mp4'
        picurl='https://cdn.download.ams.birds.cornell.edu/api/v2/asset/'+picid+'/mp4/1280'
    if os.path.exists(name+'\\'+name+'_'+i+'_'+picid+tail):
        print(name+'_'+i+'_'+picid+' '+tail+' already exists')
        return
    browser=webdriver.Chrome(options=option,desired_capabilities=capa)
    browser.get(url)
    # print(tail)
    # print(picurl)
    r=requests.get(picurl,headers=headers,proxies=random.choice(proxy))
    r.raise_for_status()
    f=open(name+'\\'+name+'_'+i+'_'+picid+tail,'wb')
    # 
    f.write(r.content)
    # i=i+1
    print(name+' '+type+'  ''+1')
    f.close()
    #['URL','ID','Name','SCIname','Auth','AgeSex','Date','Locate','Coordinates']
    ls=[]
    index=0
    a=browser.find_element(By.CSS_SELECTOR,'div.header')
    #print(a.text)
        #['URL','ID','Name','SCIname','Auth','AgeSex','Date','Locate','Coordinates']
    ls.append(name+'_'+i+'_'+picid+tail)
    ls.append(url)
    ls.append(a.find_element(By.CSS_SELECTOR,'h1.Heading.u-margine-none.Heading--minor.Heading--h5').text.replace('\n',' '))
    ls.append(a.find_element(By.CSS_SELECTOR,'span.Heading-main').text)
    ls.append(a.find_element(By.CSS_SELECTOR,'span.Heading-sub.Heading-sub--inline.Heading-sub--sci').text)
    a=browser.find_element(By.CSS_SELECTOR,'div.GridFlex.u-stack-lg.GridFlex--withGutterLarge')
    ls.append(a.find_element(By.CSS_SELECTOR,'span.main').text)
    ls.append(a.find_element(By.CSS_SELECTOR,'dd').text)
    try:
        ls.append(a.find_elements(By.CSS_SELECTOR,'span.main')[1].text)#time
    except:
        ls.append(' ')
    try:
        loc1=a.find_elements(By.CSS_SELECTOR,'span.main')[2].text
    except:
        loc1=' '
    try:
        loc2=a.find_element(By.CSS_SELECTOR,'div.u-text-4.u-stack-xs').text
    except:
        loc2=' '
    loc=loc1+' '+loc2
    ls.append(loc)

    try:
        a=browser.find_element(By.CSS_SELECTOR,'div.u-text-2.u-stack-lg')
        ls.append(a.find_elements(By.CSS_SELECTOR,'span')[-1].text)
    except:
        ls.append(' ')
    browser.close()
    browser.quit()
    f=open(name+'\\'+name+'_'+i+'_'+picid+'.txt','w',encoding='utf-8')
    for item in ls:
         f.write(item+'\n')
    f.close()
    # wb=load_workbook(name+"//"+name+'.xlsx')
    # ws=wb.active
    # ws.append(ls)
    #wb.save(name+"//"+name+'.xlsx')
    #wb.close()
      
if __name__ == "__main__":
    multiprocessing.freeze_support()
    f = open("macaulaylibrary.txt",encoding='utf-8')
    line = f.readline()
    ls=[]
    # errorls=[]
    while line:
        ls.append(line.replace('\n',''))
        line = f.readline()
    f.close()
    for name in ls:
        for item in specie_type_ls:
            savedata(name)

    # print('请输入图片，语音，视频下载的数目，以空格分开，比如 20 10 0，按回车确认')
    print('请选择下载的内容： 1.图片 2.音频 3.视频 4.全部')
    type=int(input())
    print('请输入下载数目,输入 a 下载所有：')
    num=input()
    if num=='a':
        num=999999999
    else:
        num=int(num)
    # tmp=input().split(' ')
    # try:
    #     picnum=int(tmp[0])
    #     audionum=int(tmp[1])
    #     videonum=int(tmp[2])
    # except:
    #     print("请重新输入数目：")
    #     tmp=input().split(' ')
    #     picnum=int(tmp[0])
    #     audionum=int(tmp[1])
    #     videonum=int(tmp[2])
          
    browser=webdriver.Chrome(options=option,desired_capabilities=capa)
    browser.implicitly_wait(6)
    for name in ls:
        if type!=4:
            download_species(name,num,specie_type_ls[type-1])
            sleep(5)
            savedata(name)
        else:
            download_species(name,num,specie_type_ls[0])
            sleep(5)
            savedata(name)
            download_species(name,num,specie_type_ls[1])
            sleep(5)
            savedata(name)
            download_species(name,num,specie_type_ls[2])
            sleep(5)
            savedata(name)
        # if audionum!=0:
        #     download_species(name,audionum,specie_type_ls[1])
        #     sleep(5)
        #     savedata(name,specie_type_ls[1])
        # if videonum!=0:
        #     download_species(name,videonum,specie_type_ls[2])
        #     sleep(5)
        #     savedata(name,specie_type_ls[2])

    os.system ("pause")