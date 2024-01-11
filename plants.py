import requests
import os
import time
import re
import json
from openpyxl import Workbook
from openpyxl import load_workbook

headers={
        'user-agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_14_6) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/78.0.3904.97 Safari/537.36'
        }
baseurl='https://plantsservices.sc.egov.usda.gov/api/PlantProfile?symbol='
def getone():
    wb=load_workbook('数据采集.xlsx')
    ws=wb.active
    num=ws.max_row
    for i in range(2,num):
        if ws.cell(row=i,column=3).value!= None:
            print(ws.cell(row=i,column=1).value+' already exists , pass')
            continue
        symbol=ws.cell(row=i,column=1).value
        url=baseurl+symbol
        print(url)
        try:
            r=requests.get(url,headers=headers)

        # print(r.text)
            masterid=re.findall('Id":(.*?),"Symbol',r.text)[0]
        except:
            print(symbol+' not found')
            continue
        # print('=============='+masterid)

        group=re.findall('"Group":"(.*)","RankId',r.text)[0]
        Durations=re.findall('"Durations":(.*?),"GrowthHabits',r.text)[0]
        Durations=Durations.replace(',','\n')
        Durations=Durations.replace('"','')

        Durations=Durations.replace('[','')
        Durations=Durations.replace(']','')
        # print(Durations)
        ws.cell(row=i,column=3).value=symbol
        ws.cell(row=i,column=4).value=group
        ws.cell(row=i,column=5).value=Durations
        GrowthHabit=re.findall('"GrowthHabits":(.*?),"NativeStatuses',r.text)[0]
        GrowthHabit=GrowthHabit.replace(',',' # ')
        GrowthHabit=GrowthHabit.replace('"','')
        GrowthHabit=GrowthHabit.replace('[','')
        GrowthHabit=GrowthHabit.replace(']','')
        # print(GrowthHabit)
        ws.cell(row=i,column=6).value=GrowthHabit
        NativeStatuses=re.findall('"NativeStatuses":(.*?),"MapCoordinates',r.text)[0]
        reg=re.findall('Region":"(.*?)","Status',NativeStatuses)
        status=re.findall('Status":"(.*?)","Type',NativeStatuses)
        # print(len(reg))
        NativeStatuses=''
        for k in range(0,len(reg)):
            # print(reg[i]+' '+status[i])
            # print('=====')
            NativeStatuses+=reg[k]+' '+status[k]
            if len(reg)>1:
                NativeStatuses+=' # '
        ws.cell(row=i,column=7).value=NativeStatuses
        wb.save('数据采集.xlsx')
        keydick={
            "Field": "Symbol",
            "MasterId": masterid,
            "Offset": "",
            "SortBy": "sortSciName",
            "Text": symbol
        }
        wb1=Workbook()
        ws1=wb1.active
        ws1.cell(row=1,column=1).value='Distribution Data'
        ws1.cell(row=2,column=1).value='Symbol'
        ws1.cell(row=2,column=2).value='Country'
        ws1.cell(row=2,column=3).value='State'
        ws1.cell(row=2,column=4).value='State FIP'
        ws1.cell(row=2,column=5).value='County'
        ws1.cell(row=2,column=6).value='County FIP'
        try:
            a=requests.post('https://plantsservices.sc.egov.usda.gov/api/PlantProfile/getDownloadDistributionDocumentation',data=keydick)
            a=json.loads(a.text)
            # print(a['PlantResults'])
            b=a['PlantResults'][0]
        except:
            print(symbol+' not found')
            wb1.close()
        c=b['PlantsDistributionResults']
        # for item in c:
        for j in range(0,len(c)):
            # print(item)
            # ws.cell(row=i+3,column=1).value=(c[i]["Symbol"]+','+c[i]["Country"]+','+c[i]["Country"]+','+c[i]["State"]+','+c[i]["StateFIP"]+','+c[i]["County"]+','+c[i]["CountyFIP"])
            ws1.cell(row=j+3,column=1).value=c[j]["Symbol"]
            ws1.cell(row=j+3,column=2).value=c[j]["Country"]
            ws1.cell(row=j+3,column=3).value=c[j]["State"]
            ws1.cell(row=j+3,column=4).value=c[j]["StateFIP"]
            ws1.cell(row=j+3,column=5).value=c[j]["County"]
            ws1.cell(row=j+3,column=6).value=c[j]["CountyFIP"]
        wb1.save(symbol+'.xlsx')
        wb1.close()

        

if __name__ == "__main__":
    getone()
