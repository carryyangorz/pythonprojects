from openpyxl import Workbook
from openpyxl import load_workbook
wb1=load_workbook('base.xlsx')
base=wb1.active
wb2=load_workbook('new.xlsx')
new=wb2.active
numbase=base.max_row
numnew=new.max_row

for i in range (1,numnew):
    para=new.cell(row=i,column=17).value

    if para==None:
        continue
    eeprom=new.cell(row=i,column=16).value

    para=para.replace(' ','')
    paraname=para[8:]

    if paraname == "":
        continue
    if paraname=="保留":
        continue
    if paraname=="预留":
        continue
    for k in range(1,numnew):
        if k==i:
            continue
        if (new.cell(row=k,column=16).value) == eeprom:
            print(para+'  eeprom index equals to '+ new.cell(row=k,column=17).value+ str(eeprom))
    # # print(paraname)
    for j in range(1,numbase):
        basepara=base.cell(row=j,column=17).value

        if basepara == None:
            continue
        basepara=basepara.replace(' ','')
        baseparaname=basepara[7:]
        # if paraname == baseparaname:

            # baseattribute=base.cell(row=j,column=15).value
            
            # newattribute=new.cell(row=i,column=15).value
            # if baseattribute != newattribute:
            #     print("new para:"+para+""+newattribute+"not equal"+ basepara+""+baseattribute)

            # baseattribute=base.cell(row=j,column=11).value
            # # print(baseattribute)
            
            # newattribute=new.cell(row=i,column=11).value
            # # print(newattribute)
            # if baseattribute != newattribute:
            #     print("new para:"+para+""+str(newattribute)+"not equal"+ basepara+""+str(baseattribute))

