import requests
import os
import time
import re
import json
from openpyxl import Workbook
from openpyxl import load_workbook
from selenium import webdriver
from selenium.webdriver import ChromeOptions
from time import sleep
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait 
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.desired_capabilities import DesiredCapabilities
from threading import Timer
import threading
capa = DesiredCapabilities.CHROME
capa["pageLoadStrategy"] = "eager"
option=ChromeOptions()
# option.add_argument('--headless')
option.add_argument('--no-sandbox')
option.add_argument('log-level=3') #INFO = 0 WARNING = 1 LOG_ERROR = 2 LOG_FATAL = 3 default is 0
option.add_experimental_option('excludeSwitches',['enable-automation'])
# option.add_argument('--blink-settings=imagesEnabled=false')

# option.add_experimental_option("debuggerAddress", "127.0.0.1:9527")
import multiprocessing
from multiprocessing import Process,Lock
# import multiprocessing_win
headers={
        'user-agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_14_6) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/78.0.3904.97 Safari/537.36'
        }

specie_type_ls=['photo',
'audio',
'video'
]
def savedata(data):
    bookname=data[0][0]
    data.remove(data[0])
    wb=load_workbook(bookname+'.xlsx')
    ws=wb.active
    for item in data:
        ws.append(item)
    wb.save(bookname+'.xlsx')

def savedata_2(data):
    name=data[0][0]
    specie_type=data[0][1]
    data.remove(data[0])
    wb=load_workbook(name+"//"+name+"_"+specie_type+'.xlsx')
    ws=wb.active
    for item in data:
        ws.append(item)
    wb.save(name+"//"+name+"_"+specie_type+'.xlsx')
def download_species(name,num,type):
    if not os.path.exists(name+'//'):
       os.makedirs(name+'//')
    if not os.path.exists(name+'//'+type):
       os.makedirs(name+'//'+type)

    print(name+' '+type+' searching...')
    itemnum=0
    url='https://search.macaulaylibrary.org/catalog?mediaType='+type+'&sort=rating_rank_desc'
    browser.get(url)
    browser.implicitly_wait(5)
    browser.find_element(By.CSS_SELECTOR,'input.Suggest-input').send_keys(name)
    browser.implicitly_wait(5)
    browser.find_element(By.CSS_SELECTOR,'div#Suggest-suggestion-0').click()
    sleep(3)
    a=browser.find_element(By.CSS_SELECTOR,'nav.Header-list.Header-list--dropdown')
    b=a.find_element(By.CSS_SELECTOR,'a.Header-link').get_attribute('href')
    #print(b)
    taxonCode=re.findall('&taxonCode=(.*)',b)[0]
    print(taxonCode)
    searchurl='https://macaulaylibrary.org/api/v2/search?count=100&taxonCode='+taxonCode+'&sort=rating_rank_desc'
    browser.get(searchurl)
    a=browser.page_source
    print(a)
    
    
    sleep(3000)
    flag=1
    
    while(flag):
        js="var q=document.documentElement.scrollTop=100000"
        browser.execute_script(js)
        a=browser.find_element(By.CSS_SELECTOR,'div.pagination')
        a.find_element(By.CSS_SELECTOR,'button.u-margin-none').click()
        sleep(3)
        a=browser.find_elements(By.CSS_SELECTOR,'a.ResultsGallery-link')

        if len(a)==itemnum:
             num=itemnum
             break
        itemnum=len(a)
        print(itemnum)
        if len(a)>num:
              flag=0
        
    print(num)
    if not os.path.exists(name+'//'+name+'.xlsx'):
        wb=Workbook()
        ws=wb.active
        ws.append(['URL','ID','Name','SCIname','Auth','AgeSex','Date','Locate','Coordinates'])
        wb.save(name+'//'+name+'.xlsx')
    ppls=[]
    i=0
    flag=1
    while flag:
        for j in range(20):
                if i>num-1:
                    flag=0
                    break
                id=a[i].get_attribute('data-asset-id')
                data=[id,name,type,i]
                pp=Process(target=download,args=(data,))
                i=i+1
                # if i>len(bb):
                #         break
                pp.start()
        ppls.append(pp)
        for thread in ppls:
                thread.join()
    wb.save(name+'//'+name+'.xlsx')
        
def download(data):
    option.add_argument('--headless')
    picid=data[0]
    name=data[1]
    type=data[2]
    i=str(data[3])
    url='https://macaulaylibrary.org/asset/'+picid

    if os.path.exists(name+'\\'+type+'\\'+name+'_'+i+'_'+picid+'.jpg'):
        print(name+'_'+i+'_'+picid+'.jpg already exists')
        return
    browser=webdriver.Chrome(options=option,desired_capabilities=capa)
    browser.get(url)
    picurl='https://cdn.download.ams.birds.cornell.edu/api/v2/asset/'+picid +'/2400'
    r=requests.get(picurl,headers=headers)
    r.raise_for_status()
    f=open(name+'\\'+type+'\\'+name+'_'+i+'_'+picid+'.jpg','wb')
    # 
    f.write(r.content)
    # i=i+1
    print(name+' '+type+'  ''+1')
    f.close()
    #['URL','ID','Name','SCIname','Auth','AgeSex','Date','Locate','Coordinates']
    ls=[]
    a=browser.find_element(By.CSS_SELECTOR,'div.header')
    #print(a.text)
        #['URL','ID','Name','SCIname','Auth','AgeSex','Date','Locate','Coordinates']
    ls.append(url)
    ls.append(a.find_element(By.CSS_SELECTOR,'h1.Heading.u-margine-none.Heading--minor.Heading--h5').text)
    ls.append(a.find_element(By.CSS_SELECTOR,'span.Heading-main').text)
    ls.append(a.find_element(By.CSS_SELECTOR,'span.Heading-sub.Heading-sub--inline.Heading-sub--sci').text)
    a=browser.find_element(By.CSS_SELECTOR,'div.GridFlex.u-stack-lg.GridFlex--withGutterLarge')
    ls.append(a.find_element(By.CSS_SELECTOR,'span.main').text)
    ls.append(a.find_element(By.CSS_SELECTOR,'dd').text)
    try:
        ls.append(a.find_element(By.CSS_SELECTOR,'time').text)
    except:
        ls.append(' ')
    ls.append(a.find_elements(By.CSS_SELECTOR,'span.main')[2].text+' '+a.find_element(By.CSS_SELECTOR,'div.u-text-4.u-stack-xs').text )
    a=browser.find_element(By.CSS_SELECTOR,'div.u-text-2.u-stack-lg')
    ls.append(a.find_elements(By.CSS_SELECTOR,'span')[-1].text)
    browser.close()
    browser.quit()
    f=open(name+'\\'+type+'\\'+name+'_'+i+'_'+picid+'.txt','w',encoding='utf-8')
    for item in ls:
         f.write(item+'\n')
    f.close()
    # wb=load_workbook(name+"//"+name+'.xlsx')
    # ws=wb.active
    # ws.append(ls)
    # wb.save(name+"//"+name+'.xlsx')
    # wb.close()
      
if __name__ == "__main__":
    multiprocessing.freeze_support()
    browser=webdriver.Chrome(options=option,desired_capabilities=capa)
    browser.implicitly_wait(6)
    #1.  get list workbook
#     print('请输入要下载的序号，如0，1，2,输入999下载全部')
#     i=0
#     for item in countryls:
#            print(str(i)+'.'+item+'\t')
#            i=i+1
#     choice=input()
#     if(choice=='999'):
           
#         for i in range(len(countrycodels)):
#                 get(i)
#     else:
#            get(int(choice))
       #2. download
    
    # print('请输入图片，语音，视频下载的数目，以空格分开，比如 20 10 0，按回车确认')
    # tmp=input().split(' ')
    # try:
    #     picnum=int(tmp[0])
    #     audionum=int(tmp[1])
    #     videonum=int(tmp[2])
    # except:
    #     print("请重新输入数目：")
    #     tmp=input().split(' ')
    #     picnum=int(tmp[0])
    #     audionum=int(tmp[1])
    #     videonum=int(tmp[2])
          

    # f = open("macaulaylibrary.txt",encoding='utf-8')
    # line = f.readline()
    # ls=[]
    # # errorls=[]
    # while line:
    #     ls.append(line.replace('\n',''))
    #     line = f.readline()
    # f.close()
    
    # for item in ls:
    #        if picnum!=0:
    #             download_species(item,picnum,specie_type_ls[0])
    #        if audionum!=0:
    #             download_species(item,audionum,specie_type_ls[1])
    #        if videonum!=0:
    #             download_species(item,videonum,specie_type_ls[2])
    download_species('Athene noctua',200,'photo')
    os.system ("pause")