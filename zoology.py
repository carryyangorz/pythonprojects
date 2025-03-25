import requests
import os
import random
import time
import sys
from openpyxl import Workbook
from openpyxl import load_workbook
from selenium import webdriver
from selenium.webdriver import ChromeOptions
from time import sleep
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait 
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.desired_capabilities import DesiredCapabilities
from threading import Timer
from selenium.webdriver.common.proxy import Proxy, ProxyType

capa = DesiredCapabilities.CHROME
capa["pageLoadStrategy"] = "eager"
option=ChromeOptions()
option.add_argument('--headless')
option.add_argument('--no-sandbox')
option.add_argument('log-level=3') #INFO = 0 WARNING = 1 LOG_ERROR = 2 LOG_FATAL = 3 default is 0
option.add_experimental_option('excludeSwitches',['enable-automation'])

# proxy_ip = "111.1.61.47"  # 格式 -> IP:端口 或 用户名:密码@IP:端口
# proxy_port = "3128"  # 支持 http/https/socks5
option.add_argument("--disable-gpu")        # 兼容性参数
import multiprocessing
from multiprocessing import Process,Lock
headers={
        'user-agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_14_6) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/78.0.3904.97 Safari/537.36'
        }
browser=webdriver.Chrome(options=option)
browser2=webdriver.Chrome(options=option)

def getone(name):
    if os.path.exists('result.xlsx'):
        wb=load_workbook('result.xlsx')
    else:
        wb=Workbook()
        ws=wb.active
        ws.append(['搜索名','中文名','拉丁名','来源','链接'])
        wb.save('result.xlsx')
    ws=wb.active
    ll=['形态描述','生物学','国外分布','国内分别','经济意义']


    # url='https://www.biodiversity-science.net/CN/volumn/volumn_12735.shtml'
    url = 'http://zoology.especies.cn/search/wordall?offset=0&search='+name

    browser.get(url)
    sleep(3)
    a=browser.find_elements(By.CSS_SELECTOR,'div.panel.panel-default')
    for item in a:
        b=item.find_element(By.CSS_SELECTOR,'div.panel-heading')
        if name in b.text:
            chname=b.find_elements(By.CSS_SELECTOR,'font')[0].text
            enname=b.find_elements(By.CSS_SELECTOR,'font')[1].text
            # dic={
            #     '数据源':'',
            #     '形态描述':'',
            #     '生物学':'',
            #     '国外分布':'',
            #     '国内分别':'',
            #     '经济意义':''
            #     }
            ls=[]
            print('=======================')
            print(b.text)
            theurl = b.find_element(By.CSS_SELECTOR,'a').get_attribute('href')
            print(theurl)
            if 'dbb' in theurl:
                continue
            browser2.get(theurl)
            sleep(2)
            aa=browser2.find_elements(By.CSS_SELECTOR,'div.panel.panel-primary')
            ls.append(name)
            ls.append(chname)
            ls.append(enname)
            ls.append(enname)
            ls.append(theurl)
            for item in aa:
                head=item.find_element(By.CSS_SELECTOR,'div.panel-heading').text
                body=item.find_element(By.CSS_SELECTOR,'div.panel-body').text
                bb = item.find_element(By.CSS_SELECTOR,'div.panel-footer')
                foot = bb.find_element(By.CSS_SELECTOR,'span').text
                print(name+'    +1')
                ls[3] = foot
                ls.append(head+':'+body)
            print('=======================')

            ws.append(ls)
    wb.save('result.xlsx')



if __name__ == '__main__':
    multiprocessing.freeze_support()

    if not os.path.exists('index.txt'):
        f = open('index.txt', 'w')
        f.write('0')
        f.close()
        index = 0
    else:
        f = open('index.txt', 'r')
        index = int(f.readline())
        f.close()
    # index = 0
    f = open('名录.txt', 'r',encoding='UTF-8')
    names = f.readlines()
    f.close()
    while True:
        # index=1
        try:
            item = names[index]
        except:
            print('done')
            break
        item = item.replace('\n', '')
        print(item + '    ----------------------------')
        index = index + 1
        f = open('index.txt', 'w')
        f.write(str(index))
        f.close()
        getone(item)
    os.system('pause')
